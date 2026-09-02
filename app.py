import os
import glob
import shutil
import time
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, after_this_request
from flask_cors import CORS
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import zip_processor

app = Flask(__name__, static_folder='static', template_folder='templates')
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

@app.after_request
def add_no_cache_headers(response):
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response

CORS(app, resources={r"/*": {"origins": "*"}}, expose_headers=["Content-Disposition"])
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BACKUP_DIR = os.path.join(BASE_DIR, 'backups')
os.makedirs(BACKUP_DIR, exist_ok=True)

DETAILED_HEADERS = [
    'Vendor Code', 'Party Name', 'Invoice Range', 'Total Orders',
    'New', 'Cancelled', 'Shipped', 'Delivered', 'Ready to Ship',
    'PO Created', 'Others', 'Date Range', 'Warehouse', 'Processing Status'
]

MASTER_CACHE = {
    'files_mtime': {},
    'combined_data': None,
    'files_data': {}
}

def get_workspace_files():
    files = [os.path.basename(f) for f in glob.glob(os.path.join(BASE_DIR, '*.xlsx')) if not os.path.basename(f).startswith('~$')]
    return sorted(files)

def get_files_mtime_hash():
    files = get_workspace_files()
    mtimes = {}
    for f in files:
        p = os.path.join(BASE_DIR, f)
        try:
            mtimes[f] = os.path.getmtime(p)
        except:
            mtimes[f] = 0
    return mtimes

def create_backup(filename):
    src = os.path.join(BASE_DIR, filename)
    if os.path.exists(src):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        name, ext = os.path.splitext(filename)
        backup_filename = f"{name}_backup_{timestamp}{ext}"
        dst = os.path.join(BACKUP_DIR, backup_filename)
        shutil.copy2(src, dst)
        return backup_filename
    return None

def parse_block_sheet(ws, sheet_name='', source_file=''):
    items = []
    raw_rows = list(ws.iter_rows(values_only=True))
    i = 0
    item_id = 1
    while i < len(raw_rows):
        r1 = raw_rows[i]
        if r1 and r1[0]:
            party_name = str(r1[0]).strip()
            status = str(r1[1]).strip() if len(r1) > 1 and r1[1] is not None else ''
            invoice_range = ''
            if i + 1 < len(raw_rows):
                r2 = raw_rows[i + 1]
                if r2 and r2[0] is not None:
                    invoice_range = str(r2[0]).strip()
            
            platform = 'General'
            inv_upper = invoice_range.upper()
            if inv_upper.startswith('AJ27S') or 'AJIO' in sheet_name.upper() or 'REPORT' in source_file.upper():
                platform = 'AJIO'
            elif inv_upper.startswith('MY27S') or 'MYNTRA' in sheet_name.upper() or 'PARTY' in sheet_name.upper():
                platform = 'Myntra'
            elif inv_upper.startswith('FK27S') or 'FLIPKART' in sheet_name.upper() or 'FK' in source_file.upper():
                platform = 'Flipkart'
            
            items.append({
                'id': item_id,
                'party_name': party_name,
                'status': status,
                'invoice_range': invoice_range,
                'platform': platform,
                'sheet_name': sheet_name,
                'source_file': source_file
            })
            item_id += 1
            i += 2
        else:
            i += 1
    return items

def write_block_sheet(ws, items):
    """Writes 3-row-per-entry checklist data with vivid Green / Red block styling in Excel"""
    done_card_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    done_status_fill = PatternFill(start_color="34D399", end_color="34D399", fill_type="solid")
    done_status_font = Font(name="Calibri", size=10, bold=True, color="064E3B")
    done_party_font = Font(name="Calibri", size=10, bold=True, color="065F46")
    done_inv_font = Font(name="Calibri", size=10, color="047857")

    not_card_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    not_status_fill = PatternFill(start_color="F87171", end_color="F87171", fill_type="solid")
    not_status_font = Font(name="Calibri", size=10, bold=True, color="7F1D1D")
    not_party_font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    not_inv_font = Font(name="Calibri", size=10, color="B91C1C")

    pending_party_font = Font(name="Calibri", size=10, bold=True, color="111827")
    pending_inv_font = Font(name="Calibri", size=10, color="4B5563")

    border_thin = Side(border_style="thin", color="D1D5DB")
    border_done = Side(border_style="thin", color="6EE7B7")
    border_not = Side(border_style="thin", color="FCA5A5")

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    for item in items:
        party = item.get('party_name', '').strip()
        status = item.get('status', '').strip()
        inv_range = item.get('invoice_range', '').strip()
        
        if not party:
            continue
            
        r1_num = ws.max_row + (1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 0)
        c1 = ws.cell(row=r1_num, column=1, value=party)
        c2 = ws.cell(row=r1_num, column=2, value=status if status else None)
        
        r2_num = r1_num + 1
        c3 = ws.cell(row=r2_num, column=1, value=inv_range if inv_range else 'N/A')
        c4 = ws.cell(row=r2_num, column=2, value='')

        st_upper = status.upper() if status else ''
        
        if st_upper == 'DONE':
            c1.fill = done_card_fill
            c1.font = done_party_font
            c1.alignment = left_align
            c1.border = Border(left=border_done, top=border_done, right=border_done)

            c2.fill = done_status_fill
            c2.font = done_status_font
            c2.alignment = center_align
            c2.border = Border(left=border_done, top=border_done, right=border_done)

            c3.fill = done_card_fill
            c3.font = done_inv_font
            c3.alignment = left_align
            c3.border = Border(left=border_done, bottom=border_done, right=border_done)

            c4.fill = done_card_fill
            c4.border = Border(left=border_done, bottom=border_done, right=border_done)
            
        elif st_upper in ['NOT', 'FAILED']:
            c1.fill = not_card_fill
            c1.font = not_party_font
            c1.alignment = left_align
            c1.border = Border(left=border_not, top=border_not, right=border_not)

            c2.fill = not_status_fill
            c2.font = not_status_font
            c2.alignment = center_align
            c2.border = Border(left=border_not, top=border_not, right=border_not)

            c3.fill = not_card_fill
            c3.font = not_inv_font
            c3.alignment = left_align
            c3.border = Border(left=border_not, bottom=border_not, right=border_not)

            c4.fill = not_card_fill
            c4.border = Border(left=border_not, bottom=border_not, right=border_not)
            
        else:
            c1.font = pending_party_font
            c1.alignment = left_align
            c1.border = Border(left=border_thin, top=border_thin, right=border_thin)

            c2.alignment = center_align
            c2.border = Border(left=border_thin, top=border_thin, right=border_thin)

            c3.font = pending_inv_font
            c3.alignment = left_align
            c3.border = Border(left=border_thin, bottom=border_thin, right=border_thin)

            c4.border = Border(left=border_thin, bottom=border_thin, right=border_thin)
                
        r3_num = r2_num + 1
        ws.cell(row=r3_num, column=1, value='')
        
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 16

def refresh_master_cache():
    current_mtimes = get_files_mtime_hash()
    all_files = list(current_mtimes.keys())
    
    if MASTER_CACHE['combined_data'] and MASTER_CACHE['files_mtime'] == current_mtimes:
        return
    
    combined_detailed = []
    combined_block_sheets = {
        'Short List': [],
        'Party Details': [],
        'Summary': []
    }
    files_data = {}
    
    for f in all_files:
        filepath = os.path.join(BASE_DIR, f)
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            detailed_rows = []
            block_sheets = {}
            
            for sname in wb.sheetnames:
                ws = wb[sname]
                if sname == 'Detailed Summary' or (ws.max_column and ws.max_column >= 10):
                    headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
                    for row_idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                        if not any(r):
                            continue
                        row_dict = {'source_file': f, 'id': row_idx}
                        for h, val in zip(headers, r):
                            if h:
                                row_dict[h] = val if val is not None else ''
                        for dh in DETAILED_HEADERS:
                            if dh not in row_dict:
                                row_dict[dh] = ''
                        for int_col in ['Total Orders', 'New', 'Cancelled', 'Shipped', 'Delivered', 'Ready to Ship', 'PO Created', 'Others']:
                            try:
                                row_dict[int_col] = int(row_dict.get(int_col) or 0)
                            except:
                                row_dict[int_col] = 0
                        detailed_rows.append(row_dict)
                else:
                    block_sheets[sname] = parse_block_sheet(ws, sname, f)
                    
            files_data[f] = {
                'filename': f,
                'is_combined': False,
                'available_files': all_files,
                'active_file': f,
                'sheetnames': wb.sheetnames,
                'detailed_rows': detailed_rows,
                'block_sheets': block_sheets
            }
            
            if detailed_rows:
                combined_detailed.extend(detailed_rows)
            for sname, items in block_sheets.items():
                if sname in ['Short List', 'Party Details', 'Summary']:
                    combined_block_sheets[sname].extend(items)
                else:
                    if sname not in combined_block_sheets:
                        combined_block_sheets[sname] = []
                    combined_block_sheets[sname].extend(items)
        except Exception as e:
            print(f"Error parsing {f}: {e}")
            
    for sname, items in combined_block_sheets.items():
        for idx, item in enumerate(items, 1):
            item['id'] = idx
            
    for idx, r in enumerate(combined_detailed, 1):
        r['id'] = idx

    total_vendors = len(combined_detailed)
    total_orders = sum(r.get('Total Orders', 0) for r in combined_detailed)
    total_shipped = sum(r.get('Shipped', 0) for r in combined_detailed)
    total_delivered = sum(r.get('Delivered', 0) for r in combined_detailed)
    total_cancelled = sum(r.get('Cancelled', 0) for r in combined_detailed)
    total_new = sum(r.get('New', 0) for r in combined_detailed)
    total_ready_to_ship = sum(r.get('Ready to Ship', 0) for r in combined_detailed)
    total_po_created = sum(r.get('PO Created', 0) for r in combined_detailed)
    total_others = sum(r.get('Others', 0) for r in combined_detailed)

    MASTER_CACHE['files_mtime'] = current_mtimes
    MASTER_CACHE['files_data'] = files_data
    MASTER_CACHE['combined_data'] = {
        'filename': 'Combined (All 3 Files)',
        'is_combined': True,
        'available_files': all_files,
        'active_file': '__all__',
        'sheetnames': list(combined_block_sheets.keys()),
        'detailed_rows': combined_detailed,
        'block_sheets': combined_block_sheets,
        'stats': {
            'total_vendors': total_vendors,
            'total_orders': total_orders,
            'total_shipped': total_shipped,
            'total_delivered': total_delivered,
            'total_cancelled': total_cancelled,
            'total_new': total_new,
            'total_ready_to_ship': total_ready_to_ship,
            'total_po_created': total_po_created,
            'total_others': total_others
        }
    }

refresh_master_cache()

@app.after_request
def add_header(response):
    response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/data', methods=['GET'])
def get_data():
    refresh_master_cache()
    filename = request.args.get('filename', '__all__')
    
    if filename == '__all__' or not filename:
        return jsonify({'success': True, 'data': MASTER_CACHE['combined_data']})
    
    if filename in MASTER_CACHE['files_data']:
        single = MASTER_CACHE['files_data'][filename]
        det = single['detailed_rows']
        single_data = {
            'filename': filename,
            'is_combined': False,
            'available_files': get_workspace_files(),
            'active_file': filename,
            'sheetnames': single['sheetnames'],
            'detailed_rows': det,
            'block_sheets': single['block_sheets'],
            'stats': {
                'total_vendors': len(det),
                'total_orders': sum(r.get('Total Orders', 0) for r in det),
                'total_shipped': sum(r.get('Shipped', 0) for r in det),
                'total_delivered': sum(r.get('Delivered', 0) for r in det),
                'total_cancelled': sum(r.get('Cancelled', 0) for r in det),
                'total_new': sum(r.get('New', 0) for r in det),
                'total_ready_to_ship': sum(r.get('Ready to Ship', 0) for r in det),
                'total_po_created': sum(r.get('PO Created', 0) for r in det),
                'total_others': sum(r.get('Others', 0) for r in det)
            }
        }
        return jsonify({'success': True, 'data': single_data})
        
    return jsonify({'success': False, 'error': f'File not found: {filename}'}), 400

@app.route('/api/save', methods=['POST'])
def save_data():
    payload = request.get_json()
    if not payload:
        return jsonify({'success': False, 'error': 'No data provided'}), 400
    
    filename = payload.get('filename', '__all__')
    detailed_rows = payload.get('detailed_rows', [])
    block_sheets = payload.get('block_sheets', {})
    
    saved_files = []
    backups = []
    
    if filename == '__all__':
        all_files = get_workspace_files()
        
        # 1. Save Detailed Summary + Short List to Summary_Report file
        report_file = next((f for f in all_files if 'REPORT' in f.upper()), '18-08-2026-Summary_Report.xlsx')
        b1 = create_backup(report_file)
        if b1: backups.append(b1)
        
        wb_rep = openpyxl.Workbook()
        if detailed_rows:
            ws1 = wb_rep.active
            ws1.title = 'Detailed Summary'
            ws1.append(DETAILED_HEADERS)
            for r_idx, row_data in enumerate(detailed_rows, start=2):
                row_vals = [row_data.get(h, '') for h in DETAILED_HEADERS]
                ws1.append(row_vals)
        else:
            ws1 = wb_rep.active
            ws1.title = 'Short List'
            
        if 'Short List' in block_sheets:
            ws_sl = wb_rep['Short List'] if 'Short List' in wb_rep.sheetnames else wb_rep.create_sheet('Short List')
            write_block_sheet(ws_sl, block_sheets['Short List'])
            
        wb_rep.save(os.path.join(BASE_DIR, report_file))
        saved_files.append(report_file)
        
        # 2. Save Party Details to SUMMARY.xlsx
        summary_file = next((f for f in all_files if 'SUMMARY' in f.upper() and 'REPORT' not in f.upper() and 'FK' not in f.upper()), '18-08-2026-SUMMARY.xlsx')
        b2 = create_backup(summary_file)
        if b2: backups.append(b2)
        
        wb_sum = openpyxl.Workbook()
        ws_pd = wb_sum.active
        ws_pd.title = 'Party Details'
        if 'Party Details' in block_sheets:
            write_block_sheet(ws_pd, block_sheets['Party Details'])
        wb_sum.save(os.path.join(BASE_DIR, summary_file))
        saved_files.append(summary_file)
        
        # 3. Save Summary to Summary-FK.xlsx
        fk_file = next((f for f in all_files if 'FK' in f.upper()), '18-08-2026-Summary-FK.xlsx')
        b3 = create_backup(fk_file)
        if b3: backups.append(b3)
        
        wb_fk = openpyxl.Workbook()
        ws_fk = wb_fk.active
        ws_fk.title = 'Summary'
        if 'Summary' in block_sheets:
            write_block_sheet(ws_fk, block_sheets['Summary'])
        wb_fk.save(os.path.join(BASE_DIR, fk_file))
        saved_files.append(fk_file)
        
        MASTER_CACHE['combined_data'] = None
        refresh_master_cache()
        
        return jsonify({
            'success': True,
            'message': f'All 3 Excel files saved with colors! ({", ".join(saved_files)})',
            'saved_files': saved_files,
            'backups': backups
        })
    else:
        filepath = os.path.join(BASE_DIR, filename)
        b = create_backup(filename)
        wb = openpyxl.Workbook()
        
        if detailed_rows:
            ws1 = wb.active
            ws1.title = 'Detailed Summary'
            ws1.append(DETAILED_HEADERS)
            for r_idx, row_data in enumerate(detailed_rows, start=2):
                row_vals = [row_data.get(h, '') for h in DETAILED_HEADERS]
                ws1.append(row_vals)
        else:
            wb.remove(wb.active)
            
        for sname, items in block_sheets.items():
            if items:
                ws = wb.create_sheet(title=sname)
                write_block_sheet(ws, items)
                
        if len(wb.sheetnames) == 0:
            wb.create_sheet('Summary')
            
        wb.save(filepath)
        
        MASTER_CACHE['combined_data'] = None
        refresh_master_cache()
        return jsonify({
            'success': True,
            'message': f'Saved {filename} with colors successfully!',
            'backup': b,
            'filename': filename
        })

@app.route('/api/clear', methods=['POST'])
def clear_data():
    """Completely clears uploaded excel data / files after taking safe backup"""
    payload = request.get_json() or {}
    clear_files = payload.get('clear_files', True)
    
    all_files = get_workspace_files()
    backups = []
    
    # 1. Take safety backups of existing files
    for f in all_files:
        b = create_backup(f)
        if b:
            backups.append(b)
            
    # 2. Reset / remove uploaded workspace files if requested
    if clear_files:
        for f in all_files:
            filepath = os.path.join(BASE_DIR, f)
            try:
                os.remove(filepath)
            except Exception as e:
                print(f"Error removing {f}: {e}")
                
    # 3. Invalidate cache completely
    global MASTER_CACHE
    MASTER_CACHE = {
        'files_mtime': {},
        'combined_data': None,
        'files_data': {}
    }
    refresh_master_cache()
    
    return jsonify({
        'success': True,
        'message': 'All uploaded Excel data cleared! (Safely saved in backups)',
        'backups': backups
    })

@app.route('/api/upload', methods=['POST'])
def upload_files():
    uploaded_files = request.files.getlist('files')
    if not uploaded_files:
        if 'file' in request.files:
            uploaded_files = [request.files['file']]
            
    if not uploaded_files or uploaded_files[0].filename == '':
        return jsonify({'success': False, 'error': 'No file selected'}), 400
    
    saved_names = []
    for file in uploaded_files:
        if file.filename.endswith('.xlsx'):
            dst_path = os.path.join(BASE_DIR, file.filename)
            file.save(dst_path)
            saved_names.append(file.filename)
            
    MASTER_CACHE['combined_data'] = None
    refresh_master_cache()
    
    return jsonify({
        'success': True,
        'saved_files': saved_names,
        'message': f'Uploaded {len(saved_names)} file(s) successfully!'
    })

@app.route('/api/shortlist/sync', methods=['POST'])
def sync_shortlist():
    payload = request.get_json() or {}
    detailed_rows = payload.get('detailed_rows', [])
    existing_items = payload.get('existing_items', [])
    target_sheet = payload.get('sheet_name', 'Short List')
    
    status_map = {item.get('party_name', '').strip(): item.get('status', '') for item in existing_items}
    
    new_items = []
    for r in detailed_rows:
        party = r.get('Party Name', '').strip()
        inv = r.get('Invoice Range', '').strip() or 'N/A'
        if party:
            new_items.append({
                'id': len(new_items) + 1,
                'party_name': party,
                'status': status_map.get(party, ''),
                'invoice_range': inv,
                'sheet_name': target_sheet
            })
            
    return jsonify({'success': True, 'items': new_items, 'sheet_name': target_sheet})

@app.route('/api/files', methods=['GET'])
def list_files():
    files = get_workspace_files()
    backups = [os.path.basename(f) for f in glob.glob(os.path.join(BACKUP_DIR, '*.xlsx'))]
    return jsonify({
        'success': True,
        'files': files,
        'backups': sorted(backups, reverse=True)
    })

@app.route('/api/export', methods=['GET'])
def export_file():
    filename = request.args.get('filename', '__all__')
    all_files = get_workspace_files()
    if filename == '__all__' or not filename:
        filename = all_files[0] if all_files else '18-08-2026-Summary_Report.xlsx'
        
    filepath = os.path.join(BASE_DIR, filename)
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True, download_name=filename)
    return jsonify({'error': 'File not found'}), 404

# ==========================================
# 📦 ZIP BUNDLE API ENDPOINTS (AJIO, MYNTRA, FLIPKART)
# ==========================================

@app.route('/api/zip/upload', methods=['POST'])
def upload_zip_bundle():
    platform = request.form.get('platform', 'AJIO')
    zip_type = request.form.get('zip_type', 'order')
    
    file = request.files.get('file') or request.files.get('zip_file')
    if not file or file.filename == '':
        return jsonify({'success': False, 'error': 'No ZIP file selected for upload'}), 400
        
    if not file.filename.lower().endswith('.zip'):
        return jsonify({'success': False, 'error': 'Selected file must be a .zip file'}), 400
        
    res = zip_processor.handle_zip_upload(platform, zip_type, file)
    return jsonify(res)

@app.route('/api/zip/status', methods=['GET'])
def get_zip_status():
    status = zip_processor.get_all_zip_status()
    return jsonify({'success': True, 'data': status})

@app.route('/api/zip/party_files', methods=['GET'])
def get_party_zip_files():
    platform = request.args.get('platform', 'AJIO')
    party_code = request.args.get('party_code', '').strip()
    if not party_code:
        return jsonify({'success': False, 'error': 'party_code is required'}), 400
        
    bundle = zip_processor.get_party_bundle(platform, party_code)
    return jsonify({'success': True, 'data': bundle})

@app.route('/api/zip/download/<platform>/<party_code>/<category>/<path:filename>', methods=['GET'])
def download_zip_party_file(platform, party_code, category, filename):
    p = zip_processor.normalize_platform(platform)
    safe_file = os.path.basename(filename)
    target_path = os.path.join(zip_processor.ZIP_STORAGE_DIR, p, party_code, category, safe_file)
    if os.path.exists(target_path):
        return send_file(target_path, as_attachment=True, download_name=safe_file)
    return jsonify({'error': 'File not found in zip storage'}), 404

@app.route('/api/zip/download_party_bundle/<platform>/<party_code>', methods=['GET'])
def download_party_bundle_zip(platform, party_code):
    p = zip_processor.normalize_platform(platform)
    party_code = str(party_code).strip()
    
    bundle = zip_processor.get_party_bundle(p, party_code)
    if not (bundle['has_od'] or bundle['has_two_more_invoice'] or bundle['has_details'] or bundle['has_summary']):
        return jsonify({'error': f'No files found for party {party_code} on {p}'}), 404
        
    temp_zip = os.path.join(zip_processor.ZIP_STORAGE_DIR, f"{party_code}_{p}_bundle.zip")
    zip_processor.export_party_as_zip(p, party_code, temp_zip)
    
    if os.path.exists(temp_zip):
        @after_this_request
        def remove_bundle_file(response):
            try:
                if os.path.exists(temp_zip):
                    os.remove(temp_zip)
            except Exception:
                pass
            return response
        return send_file(temp_zip, as_attachment=True, download_name=f"{party_code}_{p}_files_bundle.zip")
    return jsonify({'error': 'Failed to generate party bundle zip'}), 500

@app.route('/api/zip/cleanup_expired', methods=['GET', 'POST'])
def cleanup_expired_zip_route():
    result = zip_processor.clean_expired_zip_data(force=True)
    return jsonify(result)

@app.route('/api/zip/clear', methods=['POST'])
def clear_zip_storage_route():
    payload = request.get_json() or {}
    platform = payload.get('platform')
    zip_processor.clear_zip_storage_data(platform)
    return jsonify({
        'success': True, 
        'message': f"Cleared ZIP storage for {platform if platform else 'ALL platforms'} safely (backup preserved)."
    })

# Startup cleanup of expired files (runs both in gunicorn on Render and local dev)
try:
    zip_processor.clean_expired_zip_data(force=True)
except Exception as _startup_clean_err:
    print(f"Startup zip cleanup note: {_startup_clean_err}")

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print("=" * 60)
    print(f">> Instant High-Speed Server is running on port {port}...")
    print(f">> Open http://localhost:{port} in your browser")
    print("=" * 60)
    app.run(host='0.0.0.0', port=port, debug=False)

